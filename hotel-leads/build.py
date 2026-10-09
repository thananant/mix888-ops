"""Build hotel-leads outputs from data/.

  python3 build.py

Inputs:  data/hotels.json (hotel records), data/phones*.json (contacts by id),
         lead-finder.template.html (page with __DATA__ placeholder)
Outputs: index.html (GitHub Pages), lead-finder.html (claude.ai artifact),
         bangkok-cloudbeds-hotels.xlsx / .csv
"""
import csv, glob, json, os
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
P = lambda *a: os.path.join(HERE, *a)

hotels = json.load(open(P("data", "hotels.json"), encoding="utf-8"))
contacts = {}
for f in sorted(glob.glob(P("data", "phones*.json"))):
    for c in json.load(open(f, encoding="utf-8")):
        contacts[c["id"]] = c
for h in hotels:
    c = contacts.get(h["id"], {})
    h["ph"], h["ph2"], h["em"], h["ln"], h["web"] = (c.get(k) for k in ("phone", "phone2", "email", "line", "website"))

# --- why we recommend each hotel (shown on the page and in the spreadsheet)
TOURIST = {"Phra Nakhon", "Watthana", "Khlong Toei", "Bang Rak", "Pathum Wan", "Ratchathewi", "Sathon", "Samphanthawong"}
def fmt(n): return f"{n:,.0f}"
def reasons(h):
    out = ["ใช้ Cloudbeds อยู่แล้ว (เจอหน้าจองตรงบน Cloudbeds) เปิดรับระบบออนไลน์และจ่ายค่าซอฟต์แวร์อยู่"]
    no = h.get("no") or ""
    if "เครือ" in no or "เจ้าของเดียวกัน" in no:
        out.append("มีหลายสาขาหรือมีโรงแรมพี่น้อง คุยครั้งเดียวมีโอกาสได้หลายแห่ง")
    rm = h.get("rm")
    if rm:
        out.append(f"ขนาดใหญ่ {fmt(rm)} ห้อง งานหน้าบ้านและการจองเยอะ" if rm >= 50 else
                   f"ขนาดกลาง {fmt(rm)} ห้อง กำลังโต" if rm >= 20 else
                   f"ขนาดเล็ก {fmt(rm)} ห้อง คุยกับเจ้าของได้ตรง ตัดสินใจเร็ว")
    if h.get("s") and h["s"] >= 4: out.append(f"ระดับ {h['s']:g} ดาว มีงบลงทุนระบบ")
    rc, rv = h.get("rc"), h.get("rv")
    if rc and rc >= 1000: out.append(f"รีวิว {fmt(rc)} รายการ แขกเข้าพักต่อเนื่อง ธุรกิจมีรายได้สม่ำเสมอ")
    if rv is not None:
        if rv >= 9: out.append(f"คะแนนรีวิว {rv:.1f} ดีเยี่ยม ใส่ใจคุณภาพ พร้อมลงทุนต่อ")
        elif rv < 7.5: out.append(f"คะแนนรีวิว {rv:.1f} ค่อนข้างต่ำ มีจุดให้ช่วยปรับปรุงบริการ")
    p = h.get("p")
    if p:
        if 800 <= p <= 2500: out.append(f"ราคา ~{fmt(p)} ฿/คืน กลุ่มราคากลาง ตรงกลุ่มเป้าหมาย")
        elif p > 2500: out.append(f"ราคา ~{fmt(p)} ฿/คืน กลุ่มพรีเมียม รายได้ต่อห้องสูง")
    if h["ken"] in TOURIST: out.append(f"อยู่ย่านท่องเที่ยวหลัก ({h['kth']}) แข่งขันสูง ต้องการเครื่องมือช่วยขาย")
    elif h["ken"] == "Don Mueang": out.append("ใกล้สนามบินดอนเมือง แขกเข้าออกเร็ว ต้องจัดการห้องไว")
    if h.get("web"): out.append("มีเว็บไซต์ของตัวเอง ให้ความสำคัญกับการจองตรง")
    if h.get("ph") and h.get("em"): out.append("มีทั้งเบอร์โทรและอีเมล ติดต่อได้ทันที")
    elif h.get("ph"): out.append("มีเบอร์โทร ติดต่อได้ทันที")
    return out
for h in hotels:
    h["why"] = reasons(h)

# --- web pages
data = json.dumps(hotels, ensure_ascii=False, separators=(",", ":"))
page = open(P("lead-finder.template.html"), encoding="utf-8").read().replace("__DATA__", data)
open(P("lead-finder.html"), "w", encoding="utf-8").write(page)
head = ('<!doctype html><html lang="th"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">'
        '<style>body{margin:0;font-size:14px}[hidden]{display:none!important}</style></head><body>\n')
open(P("index.html"), "w", encoding="utf-8").write(head + page + "\n</body></html>\n")

# --- spreadsheet
TYPE = {"hotel": "โรงแรม", "boutique": "บูทีค", "apt": "เซอร์วิส อพาร์ตเมนต์", "hostel": "โฮสเทล/เกสต์เฮาส์"}
def tier(v, cuts, none="ไม่มีข้อมูล"):
    if v is None: return none
    for lim, lab in cuts:
        if v >= lim: return lab
    return cuts[-1][1]
star_t = lambda s: tier(s, [(4, "4-5 ดาว"), (3, "3-3.5 ดาว"), (0, "≤2.5 ดาว")], "ไม่ระบุ")
rev_t = lambda r: tier(r, [(9, "ดีเยี่ยม 9+"), (8, "ดีมาก 8-8.9"), (7, "ดี 7-7.9"), (0, "ต่ำกว่า 7")])
def price_t(p):
    if p is None: return "ไม่มีข้อมูล"
    return "≤800" if p <= 800 else "801-1,500" if p <= 1500 else "1,501-2,500" if p <= 2500 else ">2,500"

H = ["เขต", "ชื่อโรงแรม", "ประเภท", "เบอร์โทร", "เบอร์โทร 2", "อีเมล", "เว็บไซต์", "ดาว", "ระดับดาว",
     "คะแนนรีวิว (/10)", "จำนวนรีวิว", "แหล่งรีวิว", "ราคาเริ่มต้น/คืน (฿ ประมาณ)", "ช่วงราคา", "จำนวนห้อง",
     "ที่อยู่", "ลิงก์จอง Cloudbeds", "เหตุผลที่เสนอ", "หมายเหตุ"]
rows = sorted(([h["kth"], h["n"], TYPE[h["t"]], h["ph"], h["ph2"], h["em"], h["web"], h["s"], star_t(h["s"]),
                h["rv"], h["rc"], h["src"], h["p"], price_t(h["p"]), h["rm"], h["a"], h["u"], " / ".join(h["why"]), h["no"]] for h in hotels),
              key=lambda r: (r[0], -(r[9] or 0)))

wb = Workbook()
hf, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1C2457")
def sheet(ws, head, data, widths):
    ws.append(head)
    for c in ws[1]:
        c.font, c.fill, c.alignment = hf, fill, Alignment(wrap_text=True, vertical="center")
    for d in data: ws.append(d)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
ws = wb.active; ws.title = "รายชื่อทั้งหมด"
sheet(ws, H, rows, [16, 34, 13, 14, 14, 30, 30, 7, 11, 10, 10, 14, 13, 12, 9, 45, 48, 70, 45])
for row in ws.iter_rows(min_row=2):
    for c in (row[6], row[16]):
        if c.value: c.hyperlink, c.font = c.value, Font(color="2D45A0", underline="single")

g = defaultdict(list)
for r in rows: g[r[0]].append(r)
summ = []
for k, v in sorted(g.items(), key=lambda kv: -len(kv[1])):
    sc = [x[9] for x in v if x[9]]; pr = [x[12] for x in v if x[12]]
    hs = sum(1 for x in v if x[2] == TYPE["hostel"])
    summ.append([k, len(v), hs, len(v) - hs, sum(1 for x in v if x[3]),
                 round(sum(sc) / len(sc), 2) if sc else None, min(pr) if pr else None, max(pr) if pr else None])
summ.append(["รวม", len(rows), sum(x[2] for x in summ), sum(x[3] for x in summ), sum(x[4] for x in summ), None, None, None])
sheet(wb.create_sheet("สรุปตามเขต"), ["เขต", "จำนวน", "โฮสเทล", "โรงแรม/บูทีค/apt", "มีเบอร์โทร", "รีวิวเฉลี่ย", "ราคาต่ำสุด ฿", "ราคาสูงสุด ฿"],
      summ, [18, 9, 9, 16, 11, 11, 12, 12])
for title, key, order in [("แบ่งตามดาว", lambda r: r[8], ["4-5 ดาว", "3-3.5 ดาว", "≤2.5 ดาว", "ไม่ระบุ"]),
                          ("แบ่งตามรีวิว", lambda r: rev_t(r[9]), ["ดีเยี่ยม 9+", "ดีมาก 8-8.9", "ดี 7-7.9", "ต่ำกว่า 7", "ไม่มีข้อมูล"]),
                          ("แบ่งตามราคา", lambda r: r[13], ["≤800", "801-1,500", "1,501-2,500", ">2,500", "ไม่มีข้อมูล"])]:
    data = [[o, r[0], r[1], r[2], r[3], r[7], r[9], r[10], r[12], r[16]]
            for o in order for r in sorted((r for r in rows if key(r) == o), key=lambda r: (r[0], r[1]))]
    sheet(wb.create_sheet(title), ["กลุ่ม", "เขต", "ชื่อ", "ประเภท", "เบอร์โทร", "ดาว", "รีวิว", "จำนวนรีวิว", "ราคาเริ่มต้น ฿", "ลิงก์"],
          data, [14, 16, 34, 13, 14, 7, 8, 10, 12, 48])
wb.save(P("bangkok-cloudbeds-hotels.xlsx"))
with open(P("bangkok-cloudbeds-hotels.csv"), "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f); w.writerow(H); w.writerows(rows)
print(f"{len(rows)} hotels, {sum(1 for h in hotels if h['ph'])} with phone")
